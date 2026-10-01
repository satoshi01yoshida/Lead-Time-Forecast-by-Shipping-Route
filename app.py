from __future__ import annotations
from datetime import date
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from pathlib import Path
import warnings
import numpy as np
import pandas as pd
import plotly.express as px
import plotly.graph_objects as go
import streamlit as st
from prophet import Prophet
from prophet.utilities import regressor_coefficients
from sklearn.metrics import mean_absolute_error, mean_squared_error

warnings.filterwarnings("ignore")
st.set_page_config(page_title="航路別輸送リードタイム予測", page_icon="🚢", layout="wide")
REGRESSORS=["route_volume","scfi"]
ALIASES={
 "consignee":["consignee","荷受人","納入先","得意先"],"route":["route","航路","route_group","エリア"],"port":["port","港","港名","port_name"],
 "origin":["origin_port","origin","出発港","積港","pol"],"destination":["destination_port","destination","到着港","揚港","pod"],
 "departure":["originportetd","origin_port_etd","departure_date","departure","出発日","出港日","etd","atd"],
 "arrival_actual":["arrival_date","arrivaldate","到着日","入港日","ata"],
 "arrival_plan":["arrival_date_plan","arrivaldateplan","arrival_plan_date","planned_arrival_date","到着計画日","計画到着日"],
 "leadtime":["lead_time_days","lead_time","leadtime","輸送リードタイム","リードタイム","transit_time"],
 "date":["datetime","date","ds","日付","年月日","取得日","week_date"],"container":["container","コンテナ","port_volume","volume","traffic","y","yhat"],
 "scfi":["scfi","scfi_index","index","指数","今回値","scfi composite index"]}

def norm(v):return "".join(str(v).strip().lower().replace("　"," ").split()).replace("_","").replace("-","")
def col(df,key,required=True):
 lookup={norm(c):c for c in df.columns}
 for a in ALIASES[key]:
  if norm(a) in lookup:return lookup[norm(a)]
 if required:raise KeyError(f"{key}列を認識できません。実際の列: {df.columns.tolist()}")
 return None
def clean(s):return s.astype("string").str.strip().str.upper().str.replace(r"\s+"," ",regex=True)
def read_file(uploaded,default,all_sheets=False):
 if uploaded is None:
  p=Path(default)
  if not p.exists():raise FileNotFoundError(default)
  name,data=p.name,p.read_bytes()
 else:name,data=uploaded.name,uploaded.getvalue()
 ext=Path(name).suffix.lower()
 if ext in {".xlsx",".xlsm",".xls"}:return pd.read_excel(BytesIO(data),sheet_name=None if all_sheets else 0,engine="openpyxl")
 if ext==".csv":
  for enc in ("utf-8-sig","utf-8","cp932"):
   try:return pd.read_csv(BytesIO(data),encoding=enc)
   except UnicodeDecodeError:pass
 raise ValueError(f"未対応形式: {ext}")
def prep_master(sheets):
 cr,rp=[],[]
 for df in sheets.values():
  df=df.dropna(how="all");rc=col(df,"route",False)
  if rc is None:continue
  cc,pc=col(df,"consignee",False),col(df,"port",False)
  if cc is not None:
   x=df[[cc,rc]].rename(columns={cc:"consignee",rc:"route"});x.consignee=clean(x.consignee);x.route=x.route.astype("string").str.strip();cr.append(x.dropna().drop_duplicates())
  if pc is not None:
   x=df[[rc,pc]].rename(columns={rc:"route",pc:"port"});x.route=x.route.astype("string").str.strip();x.port=clean(x.port);rp.append(x.dropna().drop_duplicates())
 if not cr or not rp:raise KeyError("Route_Master.xlsxにはconsignee-route表とroute-port表が必要です。")
 return pd.concat(cr).drop_duplicates(),pd.concat(rp).drop_duplicates()
def prep_transport(raw,cr):
 cc,oc,dc,dt=col(raw,"consignee"),col(raw,"origin"),col(raw,"destination"),col(raw,"departure")
 lc,aa,ap=col(raw,"leadtime",False),col(raw,"arrival_actual",False),col(raw,"arrival_plan",False)
 if ap is None:raise KeyError("輸送データにarrival_date_plan（計画到着日）がありません。")
 x=pd.DataFrame({"container_row":np.arange(1,len(raw)+1),"consignee":clean(raw[cc]),"origin_port":clean(raw[oc]),"destination_port":clean(raw[dc]),"ds":pd.to_datetime(raw[dt],errors="coerce").dt.normalize(),"arrival_plan":pd.to_datetime(raw[ap],errors="coerce").dt.normalize()})
 if lc is not None:x["y"]=pd.to_numeric(raw[lc],errors="coerce")
 elif aa is not None:x["y"]=(pd.to_datetime(raw[aa],errors="coerce").dt.normalize()-x.ds).dt.total_seconds()/86400
 else:x["y"]=np.nan
 return x.merge(cr,on="consignee",how="left",validate="many_to_one")
def prep_port(raw,rp):
 dc,pc,vc=col(raw,"date"),col(raw,"port"),col(raw,"container")
 x=pd.DataFrame({"ds":pd.to_datetime(raw[dc],errors="coerce").dt.normalize(),"port":clean(raw[pc]),"v":pd.to_numeric(raw[vc],errors="coerce")}).dropna().groupby(["ds","port"],as_index=False).v.sum().merge(rp,on="port",how="inner")
 out=x.groupby(["route","ds"],as_index=False).agg(route_volume=("v","sum"),available_ports=("port","nunique"));exp=rp.groupby("route").port.nunique().rename("expected_ports");out=out.merge(exp,on="route");out["port_coverage"]=out.available_ports/out.expected_ports;return out
def prep_scfi(raw):
 dc,vc=col(raw,"date"),col(raw,"scfi");return pd.DataFrame({"scfi_date":pd.to_datetime(raw[dc],errors="coerce").dt.normalize(),"scfi":pd.to_numeric(raw[vc],errors="coerce")}).dropna().sort_values("scfi_date").drop_duplicates("scfi_date",keep="last")
def volume_for(route_daily,route,dates):
 h=route_daily.loc[route_daily.route==route,["ds","route_volume"]].sort_values("ds").drop_duplicates("ds")
 if len(h)<20:raise ValueError(f"{route}のPortWatch履歴不足: {len(h)}日")
 cal=pd.DataFrame({"ds":pd.date_range(min(h.ds.min(),dates.min()),max(h.ds.max(),dates.max()),freq="D")}).merge(h,on="ds",how="left");past=cal.ds<=h.ds.max();cal.loc[past,"route_volume"]=cal.loc[past,"route_volume"].interpolate(limit_direction="both");fut=cal.ds>h.ds.max()
 if fut.any():
  m=Prophet(yearly_seasonality=(h.ds.max()-h.ds.min()).days>=730,weekly_seasonality=True,daily_seasonality=False).fit(h.rename(columns={"route_volume":"y"})[["ds","y"]]);cal.loc[fut,"route_volume"]=np.maximum(m.predict(cal.loc[fut,["ds"]]).yhat.to_numpy(),0)
 return cal[cal.ds.isin(dates)][["ds","route_volume"]]
def features(base,route_daily,scfi,route):
 out=base.merge(volume_for(route_daily,route,pd.DatetimeIndex(base.ds)),on="ds",how="left",validate="many_to_one").sort_values("ds")
 return pd.merge_asof(out,scfi,left_on="ds",right_on="scfi_date",direction="nearest")
def new_model(train,width,cps):
 m=Prophet(yearly_seasonality=(train.ds.max()-train.ds.min()).days>=730,weekly_seasonality=False,daily_seasonality=False,interval_width=width,changepoint_prior_scale=cps)
 for r in REGRESSORS:m.add_regressor(r,prior_scale=5,standardize=True,mode="additive")
 return m
def fit(train,width,cps):return new_model(train,width,cps).fit(train[["ds","y",*REGRESSORS]])
def evaluate(h,width,cps,minimum):
 dates=np.array(sorted(h.ds.unique()));cut=max(1,int(len(dates)*.8))
 if cut>=len(dates):return None
 tr,te=h[h.ds<=dates[cut-1]],h[h.ds>dates[cut-1]]
 if len(tr)<minimum or te.empty:return None
 p=fit(tr,width,cps).predict(te[["ds",*REGRESSORS]]).yhat.to_numpy();y=te.y.to_numpy();e=p-y
 return {"MAE":mean_absolute_error(y,p),"RMSE":mean_squared_error(y,p)**.5,"Bias":e.mean(),"±3日以内率":np.mean(abs(e)<=3)}
def factors(h,m):
 corr=pd.DataFrame([{"要因":"航路交通量","Pearson":h.y.corr(h.route_volume),"Spearman":h.y.corr(h.route_volume,method="spearman")},{"要因":"SCFI","Pearson":h.y.corr(h.scfi),"Spearman":h.y.corr(h.scfi,method="spearman")}]);co=regressor_coefficients(m);co["要因"]=co.regressor.map({"route_volume":"航路交通量","scfi":"SCFI"});std=h[REGRESSORS].std(ddof=0);co["標準偏差影響日数"]=co.apply(lambda r:r.coef*std[r.regressor],axis=1);co["寄与率"]=(co.標準偏差影響日数.abs()/co.標準偏差影響日数.abs().sum()).fillna(0);return corr,co
def add_alert_causes(plans, coefficient_df, training_df):
    """各計画にProphet係数ベースの要因影響と主因を付与します。"""
    coefficient_map = coefficient_df.set_index("regressor")["coef"].to_dict()
    means = training_df[REGRESSORS].mean()
    output = plans.copy()
    output["route_volume_impact_days"] = (
        output["route_volume"] - means["route_volume"]
    ) * coefficient_map.get("route_volume", 0.0)
    output["scfi_impact_days"] = (
        output["scfi"] - means["scfi"]
    ) * coefficient_map.get("scfi", 0.0)

    def choose_main_cause(row):
        impacts = {
            "航路交通量": row["route_volume_impact_days"],
            "SCFI": row["scfi_impact_days"],
        }
        positive = {k: v for k, v in impacts.items() if pd.notna(v) and v > 0}
        if positive:
            return max(positive, key=positive.get)
        return "トレンド・未説明要因"

    output["main_cause"] = output.apply(choose_main_cause, axis=1)
    output["cause_explanation"] = output.apply(
        lambda row: (
            f"航路交通量 {row['route_volume_impact_days']:+.1f}日、"
            f"SCFI {row['scfi_impact_days']:+.1f}日。主因: {row['main_cause']}"
        ),
        axis=1,
    )
    return output


def make_alert_excel(plans, score, coefficient_df, correlation_df):
    """Summary、Alert Detail、Factor AnalysisのExcelレポートを生成します。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    detail = wb.create_sheet("Alert Detail")
    factor = wb.create_sheet("Factor Analysis")

    dark_blue = "1F4E78"
    teal = "0F6B78"
    green = "D9EAD3"
    yellow = "FFF2CC"
    red = "F4CCCC"
    gray = "E7E6E6"
    white = "FFFFFF"
    thin_gray = Side(style="thin", color="B7B7B7")

    ws.sheet_view.showGridLines = False
    detail.sheet_view.showGridLines = False
    factor.sheet_view.showGridLines = False

    ws.merge_cells("A1:D1")
    ws["A1"] = "輸送遅延アラートレポート"
    ws["A1"].fill = PatternFill("solid", fgColor=dark_blue)
    ws["A1"].font = Font(color=white, bold=True, size=15)
    ws["A1"].alignment = Alignment(horizontal="center")
    ws.append(["生成日時", pd.Timestamp.now().strftime("%Y-%m-%d %H:%M:%S")])
    total = len(plans)
    green_count = int((plans["delay_days"] <= 0).sum()) if total else 0
    yellow_count = int(plans["delay_days"].between(1, 7).sum()) if total else 0
    red_count = int((plans["delay_days"] >= 8).sum()) if total else 0
    max_delay = int(plans["delay_days"].max()) if total else 0
    summary_rows = [
        ["総計画件数", total],
        ["計画内件数", green_count],
        ["注意件数", yellow_count],
        ["警告件数", red_count],
        ["最大遅延見込み（日）", max_delay],
    ]
    for row in summary_rows:
        ws.append(row)
    ws.append([])
    ws.append(["モデル評価指標", "値"])
    if score:
        ws.append(["MAE（日）", score.get("MAE")])
        ws.append(["RMSE（日）", score.get("RMSE")])
        ws.append(["Bias（日）", score.get("Bias")])
        ws.append(["±3日以内率", score.get("±3日以内率")])
    else:
        ws.append(["評価", "評価データ不足"])
    for cell in ws[9]:
        cell.fill = PatternFill("solid", fgColor=teal)
        cell.font = Font(color=white, bold=True)
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 24
    ws["A4"].fill = PatternFill("solid", fgColor=green)
    ws["A5"].fill = PatternFill("solid", fgColor=yellow)
    ws["A6"].fill = PatternFill("solid", fgColor=red)

    detail_columns = [
        "container_row", "consignee", "route", "origin_port", "destination_port",
        "ds", "arrival_plan", "predicted_lt", "predicted_arrival", "delay_days",
        "alert", "main_cause", "route_volume_impact_days", "scfi_impact_days",
        "cause_explanation", "route_volume", "scfi", "scfi_date",
    ]
    detail_headers = [
        "行番号", "Consignee", "航路", "出発港", "到着港", "Origin Port ETD",
        "到着計画日", "予測LT（日）", "予測到着日", "差異（日）", "アラート",
        "主因", "航路交通量影響（日）", "SCFI影響（日）", "要因説明",
        "航路交通量", "SCFI", "SCFI参照日",
    ]
    detail.append(detail_headers)
    for _, row in plans.iterrows():
        detail.append([row.get(c) for c in detail_columns])
    for cell in detail[1]:
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(color=white, bold=True)
        cell.alignment = Alignment(horizontal="center")
    if detail.max_row > 1:
        table = Table(displayName="AlertDetailTable", ref=f"A1:R{detail.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        detail.add_table(table)
    for row_idx in range(2, detail.max_row + 1):
        status = str(detail.cell(row_idx, 11).value)
        fill = green if "計画内" in status else yellow if "注意" in status else red if "警告" in status else gray
        for cell in detail[row_idx]:
            cell.fill = PatternFill("solid", fgColor=fill)
        for col_idx in [6, 7, 9, 18]:
            detail.cell(row_idx, col_idx).number_format = "yyyy-mm-dd"
        for col_idx in [8, 10, 13, 14, 16, 17]:
            detail.cell(row_idx, col_idx).number_format = "0.00"
    widths = [10, 14, 12, 18, 22, 16, 16, 14, 16, 12, 24, 20, 20, 16, 55, 16, 14, 16]
    for idx, width_value in enumerate(widths, 1):
        detail.column_dimensions[get_column_letter(idx)].width = width_value
    detail.freeze_panes = "A2"

    factor.append(["分析区分", "要因", "値1", "値2", "説明"])
    for _, row in coefficient_df.iterrows():
        factor.append([
            "Prophet回帰係数",
            row.get("要因", row.get("regressor")),
            row.get("coef"),
            row.get("標準偏差影響日数"),
            "値1=1単位増加時の係数、値2=1標準偏差増加時の影響日数",
        ])
    for _, row in correlation_df.iterrows():
        factor.append([
            "相関分析", row.get("要因"), row.get("Pearson"), row.get("Spearman"),
            "値1=Pearson、値2=Spearman",
        ])
    alert_only = plans[plans["delay_days"] > 0]
    for _, row in alert_only.iterrows():
        factor.append([
            "個別アラート要因", row.get("main_cause"),
            row.get("route_volume_impact_days"), row.get("scfi_impact_days"),
            row.get("cause_explanation"),
        ])
    for cell in factor[1]:
        cell.fill = PatternFill("solid", fgColor=dark_blue)
        cell.font = Font(color=white, bold=True)
    factor.column_dimensions["A"].width = 24
    factor.column_dimensions["B"].width = 22
    factor.column_dimensions["C"].width = 18
    factor.column_dimensions["D"].width = 18
    factor.column_dimensions["E"].width = 70
    factor.freeze_panes = "A2"

    for sheet in (ws, detail, factor):
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.value is not None:
                    cell.border = Border(bottom=thin_gray)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.getvalue()


def alert_status(delay):
 if pd.isna(delay):return "⚪ 判定不能"
 if delay<=0:return "🟢 計画内"
 if delay<=7:return "🟡 注意（1～7日遅延）"
 return "🔴 警告（8日以上遅延）"
def style_alert(row):
 status=row["アラート"]
 color="#d4edda" if "🟢" in status else "#fff3cd" if "🟡" in status else "#f8d7da" if "🔴" in status else "#eeeeee"
 return [f"background-color: {color}; color: #111111"]*len(row)
def fchart(h,r):
 f=go.Figure();f.add_trace(go.Scatter(x=r.ds,y=r.upper,line={"width":0},showlegend=False));f.add_trace(go.Scatter(x=r.ds,y=r.lower,fill="tonexty",line={"width":0},name="予測範囲"));f.add_trace(go.Scatter(x=r.ds,y=r.leadtime,mode="lines+markers",name="予測LT"));f.add_trace(go.Scatter(x=h.ds,y=h.y,mode="markers",name="実績"));f.update_layout(height=430,xaxis_title="ETD",yaxis_title="LT（日）");return f

st.title("🚢 航路別輸送リードタイム予測")
st.caption("レベル3固定・コンテナ単位。影響要因分析と到着計画日超過アラートを表示します。")
with st.sidebar:
 st.header("入力（未選択時は既定ファイル）");tf=st.file_uploader("輸送データ",type=["xlsx","xlsm","xls"]);sf=st.file_uploader("SCFI",type=["xlsx","csv"]);pf=st.file_uploader("PortWatch",type=["csv","xlsx"]);mf=st.file_uploader("Route Master",type=["xlsx","xlsm"]);minimum=st.number_input("最低学習件数",10,500,15);width=st.slider("予測区間",.5,.95,.8,.05);cps=st.select_slider("トレンド柔軟性",[.01,.03,.05,.1,.3],.05)
try:
 traw=read_file(tf,"輸送データ.xlsx");sraw=read_file(sf,"SCFI_History.xlsx");praw=read_file(pf,"PortWatch_All.csv");ms=read_file(mf,"Route_Master.xlsx",True);cr,rp=prep_master(ms);transport=prep_transport(traw,cr);scfi=prep_scfi(sraw);route_daily=prep_port(praw,rp)
except Exception as e:st.error(f"読込エラー: {e}");st.stop()
historical=transport[transport.y.between(1,365)&transport.route.notna()].copy();future_plan=transport[transport.ds.ge(pd.Timestamp(date.today()))&transport.arrival_plan.notna()&transport.route.notna()].copy()
a,b,c,d=st.columns(4);a.metric("学習実績",f"{len(historical):,}件");b.metric("未来計画",f"{len(future_plan):,}件");c.metric("航路",route_daily.route.nunique());d.metric("SCFI最新日",scfi.scfi_date.max().strftime("%Y-%m-%d"))
a,b,c,d=st.columns(4);consignee=a.selectbox("Consignee",sorted(historical.consignee.unique()),index=None);route=""
if consignee:
 q=cr.loc[cr.consignee==consignee,"route"];route=str(q.iloc[0]) if not q.empty else ""
b.text_input("航路",route,disabled=True,key=f"route_{consignee}");lanes=historical.loc[historical.consignee==consignee,["origin_port","destination_port"]].drop_duplicates() if consignee else pd.DataFrame();opts=[f"{x.origin_port} → {x.destination_port}" for x in lanes.itertuples()];lane=c.selectbox("出発港 → 到着港",opts,index=None,disabled=not bool(consignee));start=d.date_input("予測開始日",date.today());run=st.button("予測・影響分析・遅延判定",type="primary",use_container_width=True,disabled=not bool(consignee and route and lane))
if run:
 origin,destination=lane.split(" → ",1);h=historical[(historical.consignee==consignee)&(historical.origin_port==origin)&(historical.destination_port==destination)].copy();h=features(h,route_daily,scfi,route).dropna(subset=["y",*REGRESSORS]).sort_values("ds")
 if len(h)<minimum:st.error(f"有効学習件数不足: {len(h)}件 / 必要{minimum}件");st.stop()
 model=fit(h,width,cps);score=evaluate(h,width,cps,minimum);future=pd.DataFrame({"ds":pd.date_range(pd.Timestamp(start),periods=31)});future=features(future,route_daily,scfi,route);fc=model.predict(future[["ds",*REGRESSORS]]);result=future[["ds","route_volume","scfi","scfi_date"]].copy();result["leadtime"]=np.maximum(fc.yhat,1);result["lower"]=np.maximum(fc.yhat_lower,1);result["upper"]=np.maximum(fc.yhat_upper,1);result["arrival"]=result.ds+pd.to_timedelta(result.leadtime.round(),unit="D")
 plans=future_plan[(future_plan.consignee==consignee)&(future_plan.origin_port==origin)&(future_plan.destination_port==destination)].copy()
 if not plans.empty:
  plans=features(plans,route_daily,scfi,route);pfcast=model.predict(plans[["ds",*REGRESSORS]]);plans["predicted_lt"]=np.maximum(pfcast.yhat,1);plans["predicted_arrival"]=plans.ds+pd.to_timedelta(plans.predicted_lt.round(),unit="D");plans["delay_days"]=(plans.predicted_arrival-plans.arrival_plan).dt.days;plans["alert"]=plans.delay_days.map(alert_status);plans=add_alert_causes(plans,regressor_coefficients(model),h)
 corr,coef=factors(h,model);st.session_state.update(result=result,history=h,score=score,plans=plans,corr=corr,coef=coef)
if "result" in st.session_state:
 r,h,score,plans=st.session_state.result,st.session_state.history,st.session_state.score,st.session_state.plans;st.divider();st.subheader("予測結果");x=r.iloc[0];a,b,c,d=st.columns(4);a.metric("予測LT",f"{x.leadtime:.1f}日");b.metric("下限",f"{x.lower:.1f}日");c.metric("上限",f"{x.upper:.1f}日");d.metric("予測到着日",x.arrival.strftime("%Y-%m-%d"));st.plotly_chart(fchart(h,r),use_container_width=True)
 if score:
  a,b,c,d=st.columns(4);a.metric("MAE",f"{score['MAE']:.2f}日");b.metric("RMSE",f"{score['RMSE']:.2f}日");c.metric("Bias",f"{score['Bias']:+.2f}日");d.metric("±3日以内率",f"{score['±3日以内率']:.1%}")
 st.subheader("到着遅延アラート")
 if plans.empty:st.info("選択したConsignee・発着港に、未来ETDかつarrival_date_plan入力済みの計画はありません。")
 else:
  green=(plans.delay_days<=0).sum();yellow=plans.delay_days.between(1,7).sum();red=(plans.delay_days>=8).sum();a,b,c=st.columns(3);a.metric("🟢 計画内",f"{green}件");b.metric("🟡 注意",f"{yellow}件");c.metric("🔴 警告",f"{red}件")
  alerts=plans[["container_row","consignee","route","origin_port","destination_port","ds","arrival_plan","predicted_lt","predicted_arrival","delay_days","alert","main_cause","route_volume_impact_days","scfi_impact_days","cause_explanation"]].rename(columns={"container_row":"行番号","consignee":"Consignee","route":"航路","origin_port":"出発港","destination_port":"到着港","ds":"Origin Port ETD","arrival_plan":"到着計画日","predicted_lt":"予測LT（日）","predicted_arrival":"予測到着日","delay_days":"差異（日）","alert":"アラート","main_cause":"主因","route_volume_impact_days":"航路交通量影響（日）","scfi_impact_days":"SCFI影響（日）","cause_explanation":"要因説明"})
  st.dataframe(alerts.style.apply(style_alert,axis=1),use_container_width=True,hide_index=True,column_config={"Origin Port ETD":st.column_config.DateColumn(format="YYYY-MM-DD"),"到着計画日":st.column_config.DateColumn(format="YYYY-MM-DD"),"予測到着日":st.column_config.DateColumn(format="YYYY-MM-DD"),"差異（日）":st.column_config.NumberColumn(format="%+d")})
  st.download_button("アラート一覧CSV",alerts.to_csv(index=False).encode("utf-8-sig"),"arrival_delay_alerts.csv","text/csv")
  alert_report = make_alert_excel(plans, score, st.session_state.coef, st.session_state.corr)
  st.download_button(
   "📊 遅延アラートレポート（Excel）",
   data=alert_report,
   file_name=f"Arrival_Delay_Report_{pd.Timestamp.now():%Y%m%d_%H%M}.xlsx",
   mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
   use_container_width=True,
  )
  delayed_plans = plans[plans.delay_days > 0].copy()
  if not delayed_plans.empty:
   st.markdown("#### 注意・警告の主因説明")
   for _, alert_row in delayed_plans.iterrows():
    icon = "🔴" if alert_row.delay_days >= 8 else "🟡"
    st.markdown(
     f"**{icon} {alert_row.consignee} | {alert_row.origin_port} → {alert_row.destination_port} | +{int(alert_row.delay_days)}日**  "
     f"主因: **{alert_row.main_cause}**  "
     f"航路交通量: {alert_row.route_volume_impact_days:+.1f}日、SCFI: {alert_row.scfi_impact_days:+.1f}日"
    )
 st.subheader("影響要因分析");t1,t2=st.tabs(["相関","回帰係数・寄与度"])
 with t1:st.dataframe(st.session_state.corr.round(3),use_container_width=True,hide_index=True)
 with t2:
  co=st.session_state.coef.copy();co["寄与率表示"]=co.寄与率.map(lambda v:f"{v:.1%}");st.dataframe(co[["要因","coef","標準偏差影響日数","寄与率表示"]].round(3),use_container_width=True,hide_index=True);st.plotly_chart(px.bar(co,x="要因",y="標準偏差影響日数",title="要因別の標準化影響"),use_container_width=True)
with st.expander("データ品質確認"):
 coverage=route_daily.groupby("route",as_index=False).agg(開始日=("ds","min"),終了日=("ds","max"),平均港カバー率=("port_coverage","mean"));st.dataframe(coverage,use_container_width=True,hide_index=True)
